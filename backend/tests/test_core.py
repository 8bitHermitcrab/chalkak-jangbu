"""DB·웹 없이 도는 핵심 로직 테스트: 텍스트 파싱, 계정 분류, 분개, 엑셀."""
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app.domain import DocStatus, DocType, DocumentData, ExtractionResult, LineItem, PaymentMethod
from app.services.accounting import AccountingError, build_journal, reconcile_amounts, validate
from app.services.classifier import ACCOUNT_DRINK, ACCOUNT_FOOD, ACCOUNT_SUPPLIES, ACCOUNT_UTILITIES, classify_item
from app.services.excel_export import build_workbook
from app.services.extraction import ExtractionService, finalize
from app.services.extraction.text_parser import parse_text

# 고깃집에 들어오던 정육 거래명세서를 OCR로 읽은 것 같은 텍스트
STATEMENT_TEXT = """거 래 명 세 표 (공급받는자 보관용)
2025년 10월 3일
공급자 등록번호 214-86-12345
상호 한돈유통 대표 김정육
공급받는자 등록번호 110-22-33333 상호 고기굽는집
품목 수량 단가 공급가액 세액
삼겹살 10 18,000 180,000 0
목살 5 16,000 80,000 0
쌈장 2 6,500 13,000 1,300
합계금액 274,300
"""

RECEIPT_TEXT = """다이소 망원점
사업자 123-45-67890
2025-10-05 14:22
쿠킹호일 2 3,000 6,000
위생장갑 1 2,000 2,000
과세물품가액 7,273
부가세 727
합계 8,000
신용카드 승인번호 12345678
"""


def test_classifier_restaurant_items():
    assert classify_item("한돈 삼겹살 1kg")[0] == ACCOUNT_FOOD
    assert classify_item("참이슬 후레쉬")[0] == ACCOUNT_DRINK
    assert classify_item("참숯 10kg")[0] == ACCOUNT_SUPPLIES
    assert classify_item("LPG 20kg")[0] == ACCOUNT_UTILITIES
    assert classify_item("알 수 없는 것", DocType.STATEMENT) == (ACCOUNT_FOOD, False)
    assert classify_item("알 수 없는 것", DocType.RECEIPT) == (ACCOUNT_SUPPLIES, False)


def test_parse_statement_picks_supplier_not_our_store():
    r = parse_text(STATEMENT_TEXT, own_biz_no="110-22-33333")
    assert r.doc_type == DocType.STATEMENT
    assert r.vendor_name == "한돈유통"
    assert r.vendor_biz_no == "214-86-12345"
    assert r.issue_date == date(2025, 10, 3)
    assert r.total_amount == 274_300
    assert r.payment_method == PaymentMethod.CREDIT
    assert [i.name for i in r.items] == ["삼겹살", "목살", "쌈장"]
    assert r.items[0].supply_amount == 180_000 and r.items[0].quantity == 10
    assert all(i.account == ACCOUNT_FOOD for i in r.items)


def test_parse_receipt():
    r = parse_text(RECEIPT_TEXT)
    assert r.doc_type == DocType.RECEIPT
    assert r.vendor_biz_no == "123-45-67890"
    assert r.issue_date == date(2025, 10, 5)
    assert (r.supply_amount, r.vat_amount, r.total_amount) == (7_273, 727, 8_000)
    assert r.payment_method == PaymentMethod.CARD
    assert [i.name for i in r.items] == ["쿠킹호일", "위생장갑"]


def test_reconcile_fills_missing_values():
    assert reconcile_amounts(None, 1_000, 11_000)[:3] == (10_000, 1_000, 11_000)
    assert reconcile_amounts(10_000, 1_000, None)[:3] == (10_000, 1_000, 11_000)
    s, v, t, warns = reconcile_amounts(None, None, 5_000)
    assert (s, v, t) == (5_000, 0, 5_000) and warns


def test_finalize_statement_fills_supply_and_vat():
    r = finalize(parse_text(STATEMENT_TEXT, own_biz_no="110-22-33333"))
    # 표 머리글('공급가액 세액')을 금액으로 잘못 읽지 않아야 한다.
    # 합계만 있는 명세서는 합계 전체를 공급가액으로 두고, 세액을 확인하라는 경고를 남긴다.
    assert (r.supply_amount, r.vat_amount, r.total_amount) == (274_300, 0, 274_300)
    assert any("부가세" in w for w in r.warnings)


def _statement_doc() -> DocumentData:
    return DocumentData(
        id=1, doc_type=DocType.STATEMENT, vendor_name="한돈유통", issue_date=date(2025, 10, 3),
        payment_method=PaymentMethod.CREDIT, supply_amount=273_000, vat_amount=1_300, total_amount=274_300,
        items=[
            LineItem(name="삼겹살", quantity=10, unit_price=18_000, supply_amount=180_000, vat_amount=0, account=ACCOUNT_FOOD),
            LineItem(name="목살", quantity=5, unit_price=16_000, supply_amount=80_000, vat_amount=0, account=ACCOUNT_FOOD),
            LineItem(name="쌈장", quantity=2, unit_price=6_500, supply_amount=13_000, vat_amount=1_300, account=ACCOUNT_FOOD),
        ],
    )


def test_journal_statement_on_credit_is_balanced():
    lines = build_journal(_statement_doc())
    assert sum(l.debit for l in lines) == sum(l.credit for l in lines) == 274_300
    accounts = {l.account: (l.debit, l.credit) for l in lines}
    assert accounts[ACCOUNT_FOOD] == (273_000, 0)
    assert accounts["부가세대급금"] == (1_300, 0)
    assert accounts["외상매입금"] == (0, 274_300)
    assert lines[0].description == "한돈유통 거래명세서 - 삼겹살 외 2건"


def test_journal_splits_accounts_and_absorbs_rounding():
    doc = DocumentData(
        doc_type=DocType.RECEIPT, vendor_name="마트", issue_date=date(2025, 10, 5),
        payment_method=PaymentMethod.CARD, supply_amount=27_273, vat_amount=2_727, total_amount=30_000,
        items=[LineItem(name="참이슬", amount=12_000, account=ACCOUNT_DRINK),
               LineItem(name="물티슈", amount=18_000, account=ACCOUNT_SUPPLIES)],
    )
    lines = build_journal(doc)
    debit = {l.account: l.debit for l in lines if l.debit}
    assert debit[ACCOUNT_DRINK] + debit[ACCOUNT_SUPPLIES] == 27_273
    assert [l.account for l in lines if l.credit] == ["미지급금(카드)"]
    assert sum(l.debit for l in lines) == sum(l.credit for l in lines)


def test_journal_refuses_mismatched_amounts():
    doc = _statement_doc()
    doc.total_amount = 300_000
    assert any("합계" in w for w in validate(doc))
    with pytest.raises(AccountingError):
        build_journal(doc)


def test_validate_flags_item_math():
    doc = _statement_doc()
    doc.items[0].supply_amount = 190_000  # 10 × 18,000 인데 잘못 읽힘
    warns = validate(doc)
    assert any("삼겹살" in w for w in warns)


def test_excel_has_sheets_and_formulas():
    doc = _statement_doc()
    doc.status = DocStatus.CONFIRMED
    data = build_workbook([doc], build_journal(doc))
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["거래내역", "품목상세", "분개장", "집계"]
    ws = wb["거래내역"]
    assert ws["D2"].value == "한돈유통" and ws["J2"].value == 274_300
    assert ws["J3"].value == "=SUM(J2:J2)"
    assert wb["품목상세"].max_row == 5  # 헤더 + 3품목 + 합계
    # 분개: 원재료 차변, 부가세대급금 차변, 외상매입금 대변 → 3줄 + 합계
    assert wb["분개장"]["D5"].value == "=SUM(D2:D4)"
    assert wb["분개장"]["E5"].value == "=SUM(E2:E4)"


class _Boom:
    name = "boom"

    def extract(self, jpeg, doc_type_hint=None):
        raise RuntimeError("api down")


class _Fixed:
    name = "fixed"

    def extract(self, jpeg, doc_type_hint=None):
        return ExtractionResult(doc_type=DocType.RECEIPT, total_amount=1_100, vat_amount=100, extractor="fixed",
                                confidence=0.9)


def test_extraction_falls_back_and_keeps_note():
    r = ExtractionService(_Boom(), _Fixed()).run(b"img")
    assert r.extractor == "fixed" and r.supply_amount == 1_000
    assert "AI 인식에 실패" in r.warnings[0]
    r2 = ExtractionService(_Boom()).run(b"img", DocType.STATEMENT)
    assert r2.extractor == "manual" and r2.doc_type == DocType.STATEMENT
