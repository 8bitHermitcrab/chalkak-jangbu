"""장부를 엑셀(.xlsx)로 내보낸다.

예전에 손으로 엑셀에 옮겨 적던 표를 그대로 자동으로 만든다는 게 목표다.
시트 구성:
  1. 거래내역        문서 한 장 = 한 줄 (날짜·거래처·공급가액·세액·합계)
  2. 품목상세        품목 한 개 = 한 줄 (어떤 고기를 얼마에 몇 kg 샀는지)
  3. 분개장          확정된 문서의 차변/대변
  4. 집계            월별·거래처별, 계정과목별 합계
합계 행은 값이 아니라 SUM 수식으로 넣어서, 엑셀에서 줄을 고쳐도 합계가 따라 바뀐다.
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.domain import DOC_TYPE_LABEL, PAYMENT_LABEL, STATUS_LABEL, DocumentData, JournalLine
from app.services.accounting import item_supply

MONEY = "#,##0"
HEADER_FILL = PatternFill("solid", fgColor="2B2A28")
TOTAL_FILL = PatternFill("solid", fgColor="F3EEE4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="D9D3C7")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def _header(ws: Worksheet, headers: list[str], widths: list[int]) -> None:
    ws.append(headers)
    for col, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    # 인쇄할 때 가로 방향, 한 페이지 너비에 맞춤 (월말에 출력해서 보관하는 경우)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def _finish(ws: Worksheet, money_cols: list[int], sum_cols: list[int], label_col: int = 1) -> None:
    """금액 서식, 테두리, 자동 필터, SUM 합계 행."""
    last = ws.max_row
    if last >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{last}"
    total_row = last + 1
    ws.cell(row=total_row, column=label_col, value="합계").font = BOLD
    for col in sum_cols:
        letter = get_column_letter(col)
        ws.cell(row=total_row, column=col, value=f"=SUM({letter}2:{letter}{last})" if last >= 2 else 0)
    for row in ws.iter_rows(min_row=2, max_row=total_row):
        for cell in row:
            cell.border = BORDER
            if cell.column in money_cols:
                cell.number_format = MONEY
            if cell.row == total_row:
                cell.fill, cell.font = TOTAL_FILL, BOLD


def build_workbook(documents: list[DocumentData], journal: list[JournalLine],
                   period: tuple[date | None, date | None] = (None, None)) -> bytes:
    docs = sorted(documents, key=lambda d: (d.issue_date or date.min, d.id or 0))
    wb = Workbook()

    # 1. 거래내역
    ws = wb.active
    ws.title = "거래내역"
    _header(ws, ["No", "거래일자", "구분", "거래처", "사업자번호", "결제수단", "품목 요약", "공급가액", "세액", "합계", "상태", "메모"],
            [6, 12, 11, 20, 15, 10, 28, 13, 11, 13, 10, 24])
    for n, d in enumerate(docs, start=1):
        names = [i.name for i in d.items]
        summary = (names[0] + (f" 외 {len(names) - 1}건" if len(names) > 1 else "")) if names else ""
        ws.append([n, d.issue_date, DOC_TYPE_LABEL[d.doc_type], d.vendor_name, d.vendor_biz_no,
                   PAYMENT_LABEL.get(d.payment_method, "") if d.payment_method else "", summary,
                   d.supply_amount, d.vat_amount, d.total_amount, STATUS_LABEL[d.status], d.memo])
        ws.cell(row=ws.max_row, column=2).number_format = "yyyy-mm-dd"
    _finish(ws, money_cols=[8, 9, 10], sum_cols=[8, 9, 10], label_col=2)

    # 2. 품목상세
    ws = wb.create_sheet("품목상세")
    _header(ws, ["거래일자", "거래처", "품목", "규격", "수량", "단가", "공급가액", "세액", "금액", "계정과목"],
            [12, 20, 24, 10, 8, 11, 13, 11, 13, 16])
    for d in docs:
        for i in d.items:
            supply = item_supply(i, d.vat_amount > 0)
            ws.append([d.issue_date, d.vendor_name, i.name, i.spec, i.quantity, i.unit_price, supply,
                       i.vat_amount, i.amount if i.amount is not None else supply, i.account])
            ws.cell(row=ws.max_row, column=1).number_format = "yyyy-mm-dd"
    _finish(ws, money_cols=[6, 7, 8, 9], sum_cols=[7, 8, 9], label_col=1)

    # 3. 분개장
    ws = wb.create_sheet("분개장")
    _header(ws, ["일자", "전표", "계정과목", "차변", "대변", "적요"], [12, 8, 18, 13, 13, 40])
    for line in sorted(journal, key=lambda l: (l.entry_date, l.document_id or 0, -l.debit)):
        ws.append([line.entry_date, line.document_id, line.account, line.debit or None, line.credit or None,
                   line.description])
        ws.cell(row=ws.max_row, column=1).number_format = "yyyy-mm-dd"
    _finish(ws, money_cols=[4, 5], sum_cols=[4, 5], label_col=3)

    # 4. 집계
    ws = wb.create_sheet("집계")
    start, end = period
    ws.append([f"기간: {start or '처음'} ~ {end or '현재'}  ·  문서 {len(docs)}건"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["월", "거래처", "건수", "공급가액", "세액", "합계"])
    head_row = ws.max_row
    monthly: dict[tuple[str, str], list[int]] = defaultdict(lambda: [0, 0, 0, 0])
    for d in docs:
        key = (d.issue_date.strftime("%Y-%m") if d.issue_date else "날짜 없음", d.vendor_name or "거래처 미상")
        agg = monthly[key]
        agg[0] += 1
        agg[1] += d.supply_amount
        agg[2] += d.vat_amount
        agg[3] += d.total_amount
    for (month, vendor), (cnt, s, v, t) in sorted(monthly.items()):
        ws.append([month, vendor, cnt, s, v, t])

    ws.append([])
    ws.append(["계정과목", "", "품목 수", "공급가액"])
    head_row2 = ws.max_row
    by_account: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    for d in docs:
        for i in d.items:
            by_account[i.account or "미분류"][0] += 1
            by_account[i.account or "미분류"][1] += item_supply(i, d.vat_amount > 0)
    for account, (cnt, amount) in sorted(by_account.items(), key=lambda kv: -kv[1][1]):
        ws.append([account, "", cnt, amount])

    for r in (head_row, head_row2):
        for cell in ws[r]:
            if cell.value is not None:
                cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for row in ws.iter_rows(min_row=head_row + 1):
        for cell in row:
            if cell.column in (4, 5, 6) and isinstance(cell.value, int):
                cell.number_format = MONEY
    for col, width in zip("ABCDEF", (14, 22, 9, 14, 12, 14)):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
